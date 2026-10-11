import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import fitz
from docx import Document


spec = importlib.util.spec_from_file_location("word_converter", Path(__file__).resolve().parents[1] / "convert_tool.py")
converter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(converter)


class WordToPdfFallbackTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.docx"
        self.output = Path(self.temp.name) / "output.pdf"
        office = patch("docx2pdf.convert", side_effect=NotImplementedError("Microsoft Word is unavailable"))
        soffice = patch.object(converter.shutil, "which", return_value=None)
        office.start()
        soffice.start()
        self.addCleanup(office.stop)
        self.addCleanup(soffice.stop)

    def convert(self, document):
        document.save(self.source)
        converter.convert_word_to_pdf(str(self.source), str(self.output))
        self.assertTrue(self.output.read_bytes().startswith(b"%PDF-"))
        with fitz.open(self.output) as pdf:
            return " ".join(page.get_text() for page in pdf), len(pdf)

    def test_long_paragraph_and_table_without_office_engines(self):
        document = Document()
        document.add_paragraph("A long paragraph that must be wrapped without truncation. " * 15 + "PARAGRAPH_END")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "TABLE_CONTENT"
        text, _ = self.convert(document)
        self.assertIn("PARAGRAPH_END", text)
        self.assertIn("TABLE_CONTENT", text)

    def test_fallback_paginates_and_keeps_last_record(self):
        document = Document()
        for index in range(80):
            document.add_paragraph(f"Record {index + 1}: " + "Readable text for pagination. " * 8)
        document.add_paragraph("FINAL_RECORD")
        text, pages = self.convert(document)
        self.assertGreater(pages, 1)
        self.assertIn("FINAL_RECORD", text)

    def test_literal_markup_and_line_breaks_are_preserved(self):
        document = Document()
        document.add_paragraph("A & B <example>\nSECOND_LINE")
        text, _ = self.convert(document)
        self.assertIn("A & B <example>", text)
        self.assertIn("SECOND_LINE", text)

    def test_missing_default_paragraph_style_still_converts(self):
        document = Document()
        document.add_paragraph("A long paragraph that must be wrapped without truncation. " * 15 + "PARAGRAPH_END")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_CONTENT"
        document.styles["Normal"].delete()
        self.assertIsNone(document.paragraphs[0].style)
        text, _ = self.convert(document)
        self.assertIn("PARAGRAPH_END", text)
        self.assertIn("TABLE_CONTENT", text)

    def test_unnamed_paragraph_style_uses_body_text(self):
        document = Document()
        document.add_paragraph("UNNAMED_STYLE_CONTENT")
        document.styles["Normal"].name = None
        self.assertIsNone(document.paragraphs[0].style.name)
        text, _ = self.convert(document)
        self.assertIn("UNNAMED_STYLE_CONTENT", text)

    def test_rendering_failure_keeps_its_diagnostic(self):
        document = Document()
        document.add_paragraph("Content")
        document.save(self.source)
        with patch("reportlab.platypus.SimpleDocTemplate.build", side_effect=RuntimeError("rendering failed")):
            with self.assertRaisesRegex(RuntimeError, "rendering failed") as failure:
                converter.convert_word_to_pdf(str(self.source), str(self.output))
        self.assertIsInstance(failure.exception.__cause__, RuntimeError)


if __name__ == "__main__":
    unittest.main()
