import importlib.util
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import fitz
from docx import Document


spec = importlib.util.spec_from_file_location("converter", Path(__file__).resolve().parents[1] / "convert_tool.py")
converter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(converter)


class PdfEditabilityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / "source.pdf"
        self.output = Path(self.temp.name) / "output.docx"

    def scanned_page(self, document):
        native = fitz.open()
        page = native.new_page(width=595, height=842)
        page.insert_text((60, 100), "Scanned document text is editable", fontsize=20)
        page.insert_text((60, 150), "Invoice total 1200 dollars", fontsize=18)
        png = page.get_pixmap(dpi=150).tobytes("png")
        scan = document.new_page(width=595, height=842)
        scan.insert_image(scan.rect, stream=png)
        native.close()

    def assert_editable(self, expected):
        converter.convert_pdf_to_word(str(self.source), str(self.output))
        doc = Document(self.output)
        text = " ".join(paragraph.text for paragraph in doc.paragraphs)
        for value in expected:
            self.assertIn(value, text)
        self.assertFalse(doc.inline_shapes, "Page screenshots must not replace editable text")

    def test_native_text_is_real_word_text(self):
        with fitz.open() as pdf:
            pdf.new_page().insert_text((60, 100), "Native document text is editable", fontsize=18)
            pdf.save(self.source)
        self.assert_editable(["Native document text is editable"])

    def test_scanned_text_uses_ocr_not_page_images(self):
        with fitz.open() as pdf:
            self.scanned_page(pdf)
            pdf.save(self.source)
        self.assert_editable(["Scanned document text is editable", "Invoice total 1200 dollars"])

    def test_mixed_pdf_keeps_both_native_and_scanned_text(self):
        with fitz.open() as pdf:
            pdf.new_page().insert_text((60, 100), "Native first page", fontsize=18)
            self.scanned_page(pdf)
            pdf.save(self.source)
        self.assert_editable(["Native first page", "Scanned document text is editable"])

    def test_missing_ocr_data_rejects_image_only_output(self):
        with fitz.open() as pdf:
            self.scanned_page(pdf)
            pdf.save(self.source)
        with patch.object(converter, "_ocr_data_directory", side_effect=RuntimeError("OCR data is not installed")):
            with self.assertRaisesRegex(RuntimeError, "OCR data is not installed"):
                converter.convert_pdf_to_word(str(self.source), str(self.output))
        self.assertFalse(self.output.exists())

    def test_native_table_cells_remain_editable(self):
        with fitz.open() as pdf:
            page = pdf.new_page()
            for x in [60, 210, 360]:
                page.draw_line((x, 60), (x, 150))
            for y in [60, 105, 150]:
                page.draw_line((60, y), (360, y))
            for point, value in [((70, 85), "Region"), ((220, 85), "Revenue"), ((70, 130), "North"), ((220, 130), "1200")]:
                page.insert_text(point, value, fontsize=14)
            pdf.save(self.source)
        converter.convert_pdf_to_word(str(self.source), str(self.output))
        doc = Document(self.output)
        self.assertTrue(doc.tables)
        cells = " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn("North", cells)
        self.assertIn("1200", cells)

    def test_force_ocr_recovers_partial_text_layer(self):
        with fitz.open() as pdf:
            self.scanned_page(pdf)
            pdf[0].insert_text((60, 50), "Existing header " * 20, fontsize=8)
            pdf.save(self.source)
        converter.convert_pdf_to_word(str(self.source), str(self.output), "force")
        text = " ".join(paragraph.text for paragraph in Document(self.output).paragraphs)
        self.assertIn("Scanned document text is editable", text)


if __name__ == "__main__":
    unittest.main()
